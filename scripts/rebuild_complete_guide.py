"""Rebuild the public MkDocs guide from the latest LABCOL-2376 Word source.

The Word document is treated as source content only. This script extracts its
instructional text and embedded figures; it does not execute any lab command.
"""
from __future__ import annotations

import re
import shutil
from html import escape
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\joshutan\Downloads\CLMEL26 - LABCOL-2376 - Next-Gen Calling--AI-Elevate with Webex - Lab Guide (1).docx")
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets" / "source-guide"

PAGE_RANGES = {
    "index.md": (19, 30),
    "01-access-the-lab.md": (31, 77),
    "02-webex-calling-org.md": (77, 167),
    "03-local-gateway.md": (167, 253),
    "04-captions-transcripts.md": (253, 327),
    "05-recording-summaries.md": (327, 372),
}

BULLET_PARAGRAPHS = {25, 26, 260, 261, 262, 363, 364, 367}
SKIP_PARAGRAPHS = {29}
OMIT_IMAGES: set[str] = set()
SAFE_TEXT_REPLACEMENTS: dict[int, str] = {}


def clean(text: str) -> str:
    text = text.replace("\ufffd", "'")
    text = re.sub(r"\s+", " ", text).strip()
    return text.replace("<FQDN>", "&lt;FQDN&gt;")


def image_names(document: Document, paragraph) -> list[str]:
    names: list[str] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rel_id = blip.get(qn("r:embed"))
        if rel_id and rel_id in document.part.rels:
            name = Path(str(document.part.rels[rel_id].target_ref)).name
            if name == "image26.tmp":
                name = "image26.png"
            names.append(name)
    return names


def extract_images(document: Document) -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    with ZipFile(SOURCE) as archive:
        for paragraph in document.paragraphs:
            for blip in paragraph._p.xpath(".//a:blip"):
                rel_id = blip.get(qn("r:embed"))
                if not rel_id or rel_id not in document.part.rels:
                    continue
                source_name = Path(str(document.part.rels[rel_id].target_ref)).name
                output_name = "image26.png" if source_name == "image26.tmp" else source_name
                if output_name in OMIT_IMAGES:
                    continue
                member = f"word/media/{source_name}"
                with archive.open(member) as source, (ASSETS / output_name).open("wb") as target:
                    shutil.copyfileobj(source, target)


def figure(name: str, context: str) -> str:
    alt = clean(context) or "Lab guide screenshot"
    alt = escape(alt[:180], quote=True)
    number = re.search(r"\d+", name)
    label = f"Source guide screenshot {number.group(0)}" if number else "Source guide screenshot"
    return (
        '<figure class="lab-figure">\n'
        f'  <a class="screenshot-link" href="assets/source-guide/{name}">\n'
        f'    <img class="lab-screenshot" src="assets/source-guide/{name}" alt="{alt}">\n'
        '  </a>\n'
        f"  <figcaption>{label}. Use the values assigned to your own lab pod.</figcaption>\n"
        "</figure>\n"
    )


PUTTY_BLOCK = """### Generate the Local Gateway configuration

1. Continuing on **Workstation 1**, minimize all applications and find the PowerShell script named `CL_Add_Trunk.ps1` on the desktop.
2. Right-click the file and select **Run with PowerShell**.

<figure class="lab-figure">
  <img src="assets/local-gateway/01-run-trunk-script.png" alt="Right-click CL_Add_Trunk.ps1 and select Run with PowerShell">
  <figcaption>Run the trunk-creation script from the Workstation 1 desktop.</figcaption>
</figure>

3. Wait while the script executes. The PowerShell window closes automatically when it finishes.
4. Confirm that `CL_LGW_Config-Ready.txt` appears on the desktop, then double-click it to open the generated CUBE configuration.

### Confirm that the trunk was added

The PowerShell script performs the **Add Trunk** operation in Control Hub and generates the matching CUBE configuration. Do not create a second trunk manually.

1. In Control Hub, go to **Services > PSTN & Routing > Gateway Configurations > Trunk**.
2. Confirm that the new registration-based trunk, normally named `dCloud-Auto`, appears.
3. The trunk can remain **Offline** until you paste the generated configuration into the LGW and registration completes.
4. If the trunk or `CL_LGW_Config-Ready.txt` is missing, stop and check the PowerShell output before continuing.

<figure class="lab-figure">
  <img src="assets/source-guide/image29.png" alt="Run the CL Add Trunk PowerShell script">
  <figcaption>Run the supplied script; it creates the Control Hub trunk and generates the matching LGW configuration.</figcaption>
</figure>

<figure class="lab-figure">
  <img src="assets/source-guide/image30.png" alt="PowerShell trunk creation output">
  <figcaption>Wait for the automated trunk creation and configuration generation to finish.</figcaption>
</figure>

<figure class="lab-figure">
  <img src="assets/source-guide/image31.png" alt="Generated Local Gateway configuration file">
  <figcaption>Open the generated configuration file before connecting to the LGW.</figcaption>
</figure>

### Connect to the LGW with PuTTY

1. Open **PuTTY** from the Workstation 1 taskbar or desktop.
2. In **Host Name (or IP address)**, enter `198.18.133.226`.
3. Leave **Port** set to `22`, select **SSH**, and click **Open**.
4. If PuTTY displays a host-key warning, confirm it only after checking that the address matches your assigned lab pod.
5. At the login prompts, enter username `admin` and password `dCloud123!`.

    PuTTY does not display password characters while you type. Press **Enter** after the password and wait for the `LGW#` prompt.

!!! warning "Source screenshot address mismatch"
    The written source step specifies `198.18.133.226`, but the supplied configuration screenshot shows `198.18.133.227` in its PuTTY title bar. Use the address assigned to your lab pod. Do not infer the correct address from the screenshot.

### Paste and apply the generated configuration

1. In `CL_LGW_Config-Ready.txt`, press **Ctrl+A**, then **Ctrl+C** to copy the entire configuration.
2. Return to the authenticated PuTTY window and right-click once inside the terminal to paste. The commands begin running immediately.
3. Wait until the `LGW#` prompt returns. Scroll through the terminal output and check for command errors.
4. You may ignore only the final save error shown after `copy run start`; do not ignore errors from the configuration commands.

<figure class="lab-figure">
  <img src="assets/local-gateway/02-paste-lgw-config.png" alt="Generated LGW configuration selected in Notepad and pasted into an authenticated PuTTY session">
  <figcaption>Copy the complete generated configuration and paste it into the authenticated LGW session. This supplied screenshot shows a sample pod address; use your assigned address.</figcaption>
</figure>

### Verify Webex Calling registration

1. Wait **2 to 3 minutes** for the LGW to initiate registration with Webex Calling.
2. At the `LGW#` prompt, run `show sip-ua register status`.
3. In the command output, confirm that the **reg** column reports `yes`. If it does not, recheck the pasted Voice Class Tenant and SIP profile configuration before continuing.

<figure class="lab-figure">
  <img src="assets/local-gateway/03-verify-registration.png" alt="Output of show sip-ua register status with the registration status highlighted as yes">
  <figcaption>The LGW is registered when the <strong>reg</strong> column shows <code>yes</code>.</figcaption>
</figure>

Once the registration status is `yes`, continue to Module 2b.
"""


def render_page(document: Document, filename: str, start: int, end: int) -> str:
    lines: list[str] = []
    last_context = "Lab guide screenshot"
    first_heading = True
    inserted_putty = False
    ordered_counter = 0

    if filename != "index.md":
        lines.extend([
            "> **Lab data notice:** Values in screenshots are examples from the source guide. Use the credentials, addresses, phone numbers, and session details assigned to your own pod.",
            "",
        ])

    for index in range(start, min(end, len(document.paragraphs))):
        paragraph = document.paragraphs[index]
        text = SAFE_TEXT_REPLACEMENTS.get(index, clean(paragraph.text))
        images = [name for name in image_names(document, paragraph) if name not in OMIT_IMAGES]

        if filename == "03-local-gateway.md" and 185 <= index <= 209:
            if not inserted_putty:
                lines.extend([PUTTY_BLOCK.rstrip(), ""])
                inserted_putty = True
            continue
        if index in SKIP_PARAGRAPHS:
            continue

        if text:
            last_context = text
            style = paragraph.style.name
            if style == "Heading 1":
                ordered_counter = 0
                level = "#" if first_heading else "##"
                lines.extend([f"{level} {text}", ""])
                first_heading = False
            elif style == "Heading 2":
                ordered_counter = 0
                lines.extend([f"## {text}", ""])
            elif style == "Heading 3":
                ordered_counter = 0
                lines.extend([f"### {text}", ""])
            elif style == "List Paragraph":
                if index in BULLET_PARAGRAPHS:
                    marker = "-"
                    ordered_counter = 0
                else:
                    ordered_counter += 1
                    marker = f"{ordered_counter}."
                lines.append(f"{marker} {text}")
            else:
                lines.extend([text, ""])

        for name in images:
            lines.extend([figure(name, last_context).rstrip(), ""])

    title_prefix = "" if filename == "index.md" else ""
    return title_prefix + "\n".join(lines).strip() + "\n"


def add_index_front_matter(body: str) -> str:
    header = """---
title: Home
hide:
  - navigation
  - toc
---

<div class="lab-hero">
<p class="eyebrow">CISCO LIVE MELBOURNE 2026 · LABCOL-2376 · HANDS-ON LAB</p>
<h1>Next-Gen Calling<br>AI-Elevate with Webex</h1>
<p class="hero-description">Configure Webex Calling and a Local Gateway, then experience AI-powered captions, transcripts, and recording summaries.</p>
<p class="hero-authors">Hussain Ali and Omer Ilyas · Technical Marketing Engineers</p>
<a class="hero-button" href="01-access-the-lab.html">Get started <span aria-hidden="true">→</span></a>
<a class="hero-secondary" href="03-local-gateway.html">Go to the Local Gateway lab →</a>
</div>

## Your lab journey

<div class="journey-grid">
<a class="journey-card" href="01-access-the-lab.html"><span class="card-number">01</span><h3>Get ready</h3><p>Connect to the lab environment, open Workstation 1, and locate your assigned credentials.</p><span>Access the lab →</span></a>
<a class="journey-card" href="02-webex-calling-org.html"><span class="card-number">02</span><h3>Build the calling foundation</h3><p>Configure the Webex Calling organization, phone numbers, users, and the registration-based Local Gateway.</p><span>Organization and gateway →</span></a>
<a class="journey-card" href="04-captions-transcripts.html"><span class="card-number">03</span><h3>Experience Calling AI</h3><p>Test live captions, call transcripts, recording, and AI-generated recording summaries.</p><span>Captions and summaries →</span></a>
</div>
"""
    body = re.sub(r"^# Webex Calling AI Overview\s*", "", body)
    return header + body


def main() -> None:
    document = Document(SOURCE)
    extract_images(document)
    for filename, (start, end) in PAGE_RANGES.items():
        content = render_page(document, filename, start, end)
        if filename == "index.md":
            overview = re.sub(
                r"^# Webex Calling AI Overview",
                "# Lab overview",
                content,
                count=1,
            )
            (DOCS / "00-lab-overview.md").write_text(overview, encoding="utf-8")
            print("Wrote 00-lab-overview.md")
            content = add_index_front_matter("")
        (DOCS / filename).write_text(content, encoding="utf-8")
        print(f"Wrote {filename}")


if __name__ == "__main__":
    main()
